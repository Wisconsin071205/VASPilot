"""Apply positioning edits directly to the user-selected Word originals."""
from datetime import datetime
from pathlib import Path
import shutil
from zipfile import ZipFile

from docx import Document
from revise_positioning import INTRO, DESIGN, TITLE

ROOT = Path(__file__).resolve().parents[1] / 'competition_submission' / 'ok'


def set_text(paragraph, text):
    # Keep paragraph layout and all non-text runs (including anchored images).
    runs = [r for r in paragraph.runs if r.text]
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ''


def media(path):
    with ZipFile(path) as archive:
        return {n: archive.read(n) for n in archive.namelist() if n.startswith('word/media/')}


def revise(path, backup):
    shutil.copy2(path, backup / path.name)
    original_media = media(path)
    doc = Document(path)
    is_intro = '简介' in path.name
    changes = INTRO if is_intro else DESIGN
    changed = 0
    for paragraph in doc.paragraphs:
        value = paragraph.text
        if value == 'vasp 调度计算智能体':
            set_text(paragraph, TITLE)
            changed += 1
            continue
        # Feature paragraphs have a bold label followed by body text. Edit
        # these runs separately so their typographic hierarchy survives.
        if is_intro and paragraph.runs and paragraph.runs[0].bold:
            for run in paragraph.runs:
                for prefix, replacement in changes.items():
                    if run.text.startswith(prefix):
                        trailing = '  ' if run.text.endswith('  ') else ''
                        run.text = replacement + trailing
                        changed += 1
                        break
            continue
        for prefix, replacement in changes.items():
            if value.startswith(prefix):
                # Retain this source's existing architecture image and caption.
                if prefix.startswith('图 1'):
                    continue
                set_text(paragraph, replacement)
                changed += 1
                break
    if not is_intro:
        # Align the navigation with the user's deliberately shortened chapter.
        for table in doc.tables:
            for row in table.rows:
                if row.cells[0].text == '7  部署与验证设计':
                    set_text(row.cells[0].paragraphs[0], '7  部署')
                    set_text(row.cells[1].paragraphs[0], '运行环境与启动流程')
        for paragraph in doc.paragraphs:
            if paragraph.text.startswith('本机需要 Python 3.11'):
                set_text(paragraph, '本机需要 Python 3.11 或更高版本、OpenSSH 和浏览器，并能访问所配置的模型服务。目标服务器需具备 SSH 可达性及所选连接方式需要的 Python 3 辅助环境；作业管理依赖 Slurm 或 PBS，VASP 专用流程另需有权使用的 VASP、MPI 和赝势环境。Vlab 路由部署 Gateway，直连使用目标端辅助程序；VS Code Bridge 按需安装。')
            if paragraph.text.startswith('建议依次完成：'):
                set_text(paragraph, '建议依次完成：安装并启动本地控制台 → 登记多台服务器及各自工作根目录、认证和路由 → 分别诊断连接、目录与队列 → 配置模型并探测 → 指定目标开展文件与作业操作。验证 VASP 专用流程时，再使用小型算例完成校验、审阅、审批、提交和结果回收。首次运行应明确目标服务器、分区、核数与时限。')
    else:
        # Avoid an empty page when a standalone page-break paragraph spills.
        for paragraph in doc.paragraphs:
            if not paragraph.text.strip() and paragraph._p.xpath('.//w:br[@w:type="page"]'):
                for br in paragraph._p.xpath('.//w:br[@w:type="page"]'):
                    br.getparent().remove(br)
                paragraph.paragraph_format.space_before = 0
                paragraph.paragraph_format.space_after = 0
                paragraph.paragraph_format.line_spacing = 0.1
            if paragraph.text.startswith('4  模型及 API'):
                paragraph.paragraph_format.page_break_before = True
    doc.core_properties.title = TITLE + ('作品简介' if is_intro else '作品设计文档')
    doc.save(path)
    assert media(path) == original_media, 'Original image data changed'
    print(path.name, changed, 'text edits;', len(original_media), 'original images preserved')


if __name__ == '__main__':
    backup = ROOT / ('原稿备份_' + datetime.now().strftime('%Y%m%d_%H%M%S'))
    backup.mkdir(parents=True, exist_ok=False)
    for path in ROOT.glob('*.docx'):
        if not path.name.startswith('~$'):
            revise(path, backup)
    print('Backup:', backup)
