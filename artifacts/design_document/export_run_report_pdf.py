from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1] / 'competition_submission' / 'ok'
source = ROOT / '远端控制智能体_算力集群运行数据说明_50785.docx'
output = ROOT / '远端控制智能体_算力集群运行数据说明_50785.pdf'
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
styles = getSampleStyleSheet()
body = ParagraphStyle('body', parent=styles['BodyText'], fontName='STSong-Light', fontSize=10.3, leading=16, spaceAfter=6, wordWrap='CJK')
title = ParagraphStyle('title', parent=body, fontSize=21, leading=28, alignment=TA_CENTER, spaceAfter=12)
subtitle = ParagraphStyle('subtitle', parent=body, fontSize=12, leading=18, alignment=TA_CENTER, spaceAfter=8)
h1 = ParagraphStyle('h1', parent=body, fontSize=15, leading=22, spaceBefore=8, spaceAfter=8)
h2 = ParagraphStyle('h2', parent=body, fontSize=11.5, leading=18, spaceBefore=5, spaceAfter=5)
code = ParagraphStyle('code', parent=body, fontName='Courier', fontSize=9, leading=13, leftIndent=8, rightIndent=8)
cell = ParagraphStyle('cell', parent=body, fontSize=8.6, leading=12, spaceAfter=0)

def footer(canvas, doc):
    canvas.saveState(); canvas.setFont('STSong-Light', 8); canvas.setFillColor(colors.HexColor('#666666'))
    canvas.drawCentredString(A4[0] / 2, 10 * mm, f'远端控制智能体算力集群运行数据说明 · 第 {doc.page} 页'); canvas.restoreState()

def para(text, style=body):
    return Paragraph(escape(text).replace('\n', '<br/>'), style)

doc = Document(source)
story = []
paragraphs = {p._p: p for p in doc.paragraphs}
tables = {t._tbl: t for t in doc.tables}
for child in doc.element.body.iterchildren():
    tag = child.tag.rsplit('}', 1)[-1]
    if tag == 'p':
        p = paragraphs.get(child)
        if p is None or not p.text.strip(): continue
        if p.paragraph_format.page_break_before: story.append(PageBreak())
        name = p.style.name if p.style else 'Normal'
        style = title if name == 'Title' else subtitle if name == 'Subtitle' else h1 if name == 'Heading 1' else h2 if name == 'Heading 2' else code if p.text.startswith('reached required accuracy') else body
        story.append(para(p.text, style))
    elif tag == 'tbl':
        table_obj = tables.get(child)
        if table_obj is None: continue
        rows = [[para(c.text, cell) for c in row.cells] for row in table_obj.rows]
        widths = [55*mm, 75*mm] if len(rows[0]) == 2 else [32*mm, 48*mm, 50*mm]
        table = Table(rows, colWidths=widths, repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#D9D9D9')), ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EEF1F4')), ('VALIGN',(0,0),(-1,-1),'MIDDLE'), ('LEFTPADDING',(0,0),(-1,-1),5), ('RIGHTPADDING',(0,0),(-1,-1),5), ('TOPPADDING',(0,0),(-1,-1),4), ('BOTTOMPADDING',(0,0),(-1,-1),4)]))
        story.extend([Spacer(1,3), table, Spacer(1,8)])

SimpleDocTemplate(str(output), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=17*mm, bottomMargin=17*mm, title='远端控制智能体算力集群运行数据说明 作业50785').build(story, onFirstPage=footer, onLaterPages=footer)
print(output)
