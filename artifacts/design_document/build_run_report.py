from pathlib import Path
import json, hashlib, shutil, zipfile
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'competition_submission' / 'ok'
EVIDENCE = ROOT / 'run_evidence_50785'
EVIDENCE.mkdir(exist_ok=True)
inputs = [Path(r'D:\lgy\下载\slurm-evidence-50785.json'), Path(r'D:\lgy\下载\slurm-evidence-50785 (1).md'), Path(r'C:\Users\Lenovo\.codex\attachments\205b3a5b-76c7-46f1-ad6b-1e2ec9015e66\pasted-text.txt')]
for source in inputs:
    shutil.copy2(source, EVIDENCE/source.name)
data = json.loads(inputs[0].read_text(encoding='utf-8-sig'))
s = data['summary']
wait = (datetime.fromisoformat(s['start'])-datetime.fromisoformat(s['submit'])).total_seconds()
elapsed = (datetime.fromisoformat(s['end'])-datetime.fromisoformat(s['start'])).total_seconds()
assert wait == 37 and elapsed == 3402
for log in data['logs']:
    (EVIDENCE/(Path(log['path']).name+'.excerpt.txt')).write_text(log['content'],encoding='utf-8')
(EVIDENCE/'sacct.raw.txt').write_text(data['sources']['sacct']['raw'],encoding='utf-8')
hashes = {f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in EVIDENCE.iterdir() if f.is_file() and f.name != 'SHA256SUMS.txt'}
(EVIDENCE/'SHA256SUMS.txt').write_text('\n'.join(f'{v}  {k}' for k,v in sorted(hashes.items()))+'\n',encoding='utf-8')

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72)
sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
for name,size in [('Normal',11),('Title',23),('Subtitle',12),('Heading 1',16),('Heading 2',12)]:
    style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '微软雅黑')
    style.paragraph_format.space_after=Pt(7);style.paragraph_format.line_spacing=1.13
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
footer=sec.footer.paragraphs[0];footer.alignment=1
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='晶织算力集群运行数据说明 作业50785'
doc.core_properties.subject='依据用户提供的Slurm记账与日志摘录整理的正式说明'

def p(text):return doc.add_paragraph(text)
def h(text):return doc.add_heading(text,2)
def page(title):
    paragraph=doc.add_heading(title,1);paragraph.paragraph_format.page_break_before=True
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,row in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(text)
            props=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
            props.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement('w:'+edge);node.set(qn('w:w'),'80');node.set(qn('w:type'),'dxa');margins.append(node)
            props.append(margins)
            if i==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'EEF1F4');props.append(fill)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                for run in para.runs:run.font.size=Pt(10);run.bold=(i==0)
        if i==0:
            repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    p('').paragraph_format.space_after=Pt(0)

doc.add_paragraph('算力集群运行数据说明',style='Title')
p('晶织 多服务器连接与计算控制智能体')
p('作业 50785  |  正式说明  |  整理日期 2026年9月6日')
p('本报告依据用户提供的 Slurm 导出数据与运行日志摘录，说明 Fe_relax_128 作业的运行、等待和内存记账情况。该作业运行在课题组自有服务器 cl12（集群标识 cluster_mgt），由晶织提交；在 node5 分配 64 个 CPU，运行 56 分 42 秒，最终状态为 COMPLETED，退出码为 0:0。该服务器不属于比赛指定的本科生算力平台，因此本报告作为课题组服务器上的实际运行证明与作品功能证据提交，不冒充本科生算力平台证明。')
h('1  作业与环境')
table(['项目','记录值'],[
('作品实现标识','rsci'),('服务器配置名 / 集群标识','cl12 / cluster_mgt'),
('作业 ID / 名称','50785 / Fe_relax_128'),('分区 / 执行节点','normal / node5'),
('分配 CPU 数','64（Slurm AllocCPUS）'),('任务规模','日志显示 Fe，1 类元素，128 个离子'),
('软件环境','job-50785.err：vasp/5.4.4-oneapi-2024.1'),
('计算目录',data['directory']),('证据采集时间','2026-09-06 06:21:56 UTC')],[1.65,5.15])
h('2  作业执行结论')
p('sacct 主记录与 batch、extern、0 三个步骤均为 COMPLETED，退出码均为 0:0。该结论说明提供的调度记账记录中作业已成功结束，不单独代表科学计算结果已通过独立验收。JSON 与 Markdown 是同一采集快照的不同表示，不视为两次独立运行证据。')
p('数据来源：[1] JSON 中 records.sacct、summary 和 logs；[2] Markdown 导出。平台归属已由提交者确认：cl12 / cluster_mgt 为课题组自有服务器。')

page('3  排队与运行时长')
table(['事件','时间或时长'],[
('提交时间','2026-09-05 00:12:38'),('实际开始','2026-09-05 00:13:15'),('实际结束','2026-09-05 01:09:57'),
('提交至实际开始','37 秒 = Start − Submit'),('Slurm 运行时长','00:56:42 = 3402 秒'),
('OUTCAR 程序计时','Elapsed time：3363.003 秒（约 56 分 3 秒）')],[2.1,4.7])
p('上述三个事件时间均保留 Slurm 原始集群时间，记录没有时区偏移，未转换为北京时间。VASP 程序计时比 Slurm 作业生命周期短 38.997 秒；两者统计范围不同，现有材料不能将差值精确归因于环境加载或其他某项操作。')
h('3.1  排队证据范围')
p('本次在作业结束后采集，squeue 返回 Invalid job id specified，sstat 无记录。现有材料可计算提交至实际启动的等待时间，但没有 PENDING 状态快照或等待原因历史。node5 是执行节点，不能作为排队原因。历史记账记录仍可用于说明作业的起止时间。[1][4]')
h('3.2  内存占用记账')
table(['作业步骤','MaxRSS 原值','MaxVMSize 原值'],[
('50785.batch','11268K','563116K'),('50785.extern','1612K','412344K'),
('50785.0\nhydra_bstrap_proxy','22291856K','612953100K')],[2.25,2.2,2.35])
p('申请内存原值为 3948Mc，属于请求配置，不能当作实际占用。主作业行的 MaxRSS 为空；可用值来自步骤记录。50785.0 的 MaxRSS 按 1024 进制折算约为 21.26 GiB，仅用于帮助阅读，正式证据保留原始单位。')
p('Slurm 将步骤内任务的 MaxRSS 取最大值；此处不是已核实的全作业峰值总内存。不能相加三个步骤的峰值，也不能乘以 64 推算整体峰值。MaxVMSize 是虚拟内存指标，不能替代常驻内存；现有数据也不足以绘制真实内存随时间变化曲线。[4]')

page('4  运行日志与结果说明')
h('4.1  软件启动与任务规模')
p('job-50785.err 记录加载 vasp/5.4.4-oneapi-2024.1 模块；vasp.out 记录使用 64 total cores，并识别 1 types and 128 ions。日志中的 64 核与 Slurm 分配的 64 个 CPU 相互对应。提交者确认该 Fe-relax 目录在作业结束后没有运行其他任务，因此本报告将这些通用日志视为作业 50785 的运行证据。')
h('4.2  程序完成标记')
p('vasp.out 中出现如下原始结束标记：')
q=p('reached required accuracy - stopping structural energy minimisation')
for run in q.runs:run.font.name='Consolas';run.font.size=Pt(10)
p('OSZICAR 摘录列出 3 个离子步骤；最后一条为 F = −0.10552294E+04，E0 = −0.10551671E+04，mag = 279.0294。报告保留这些值作为日志摘录，不据此宣称材料性质已获得科研验证。')
h('4.3  程序内部统计')
table(['OUTCAR 字段','原始值'],[
('Total CPU time used (sec)','3198.278'),('Elapsed time (sec)','3363.003'),
('Maximum memory used (kb)','375040'),('VASP MPI-rank0 内存明细总值','130771 kBytes')],[4.1,2.7])
p('这些值属于 OUTCAR 程序内部统计，与 Slurm 步骤记账不采用同一汇总范围；本报告不将其当作整个 64 CPU 作业的累计 CPU 时间或总内存，也不据此计算并行效率。Average memory used 为 0，仅按原日志保留，不解读为运行时未使用内存。')
h('4.4  与作品功能的对应')
p('用户粘贴的界面文字将作业 50785、服务器 cl12、计算目录、自动分析及证据下载关联在一起，可说明作品提供了该作业的查询、日志分析和证据导出入口。提交者确认该作业由晶织提交；本次附件仍未包含提交脚本、审批记录或提交成功截图，因此报告以界面关联和 Slurm 结果作为证据，不虚构未提供的操作画面。')
p('界面规则引擎 builtin-rules-v1 显示“电子与离子收敛判据均已满足”。现有摘录支持程序报告结构优化停止条件已满足；本报告未用缺失的完整 INCAR 重新计算判据，但提交者确认日志未被后续任务覆盖。[1][3]')

page('5  证据清单与提交核对')
table(['附件','内容及作用'],[
('[1] slurm-evidence-50785.json','原始采集命令、Slurm 返回记录、日志摘录及解析摘要'),
('[2] slurm-evidence-50785 (1).md','同次采集的可读导出材料'),
('[3] pasted-text.txt','用户粘贴的界面文字，含作业列表和自动分析'),
('*.excerpt.txt / sacct.raw.txt','从 JSON 按原内容提取，便于检查；不是额外采集'),
('SHA256SUMS.txt','附件文件哈希，用于核对整理后文件一致性；不是平台签章')],[3,3.8])
h('5.1  提交前需确认的事项')
table(['核对事项','当前情况与所需补充'],[
('平台归属','已确认 cl12 / cluster_mgt 是课题组自有服务器，不是比赛指定的本科生算力平台；提交时按“作品实际运行证明”说明。'),
('作品参与环节','已确认作业 50785 由晶织提交；附件暂未包含提交脚本、审批记录或提交成功截图。'),
('时间字段差异','界面顶部出现 2026/09/05 13:51:48，未标明该字段含义，与 Slurm 起止时间不同；报告以 sacct 为准并保留差异说明。'),
('通用日志归属','已确认 Fe-relax 目录在作业结束后没有运行其他任务，未覆盖 vasp.out、OUTCAR、OSZICAR。'),
('排队与实时内存','本次无历史排队快照、无 sstat 采样；如没有留存，保持缺失说明，不补画历史曲线。')],[1.65,5.15])
h('5.2  结论与引用')
p('提供的 Slurm 记录能够说明课题组服务器作业 50785 的成功结束、64 CPU 分配、37 秒提交至启动等待、56 分 42 秒运行时长，以及步骤级内存记账。配套日志说明目录内存在 VASP 运行与正常停止记录；用户已确认作业由晶织提交且目录未被后续任务覆盖。本报告可作为作品在课题组 Slurm 服务器上的运行证明提交，但不能替代比赛指定本科生算力平台的证明。')
p('[4] Slurm 官方 sacct 字段说明：https://slurm.schedmd.com/sacct.html（核对日期：2026年9月6日）。本报告依据用户提供文件整理，未直接登录集群复采，也不替代平台出具的认证。')

output=OUT/'晶织_算力集群运行数据说明_50785.docx'
doc.save(output)
with zipfile.ZipFile(OUT/'50785_原始证据.zip','w',zipfile.ZIP_DEFLATED) as archive:
    for f in sorted(EVIDENCE.iterdir()):
        if f.is_file():archive.write(f,arcname=f.name)
print(output)
