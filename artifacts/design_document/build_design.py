from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent/'competition_submission'/'修订版'/'晶织_作品设计文档.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.1); sec.bottom_margin=Cm(2.0)
sec.left_margin=Cm(2.3); sec.right_margin=Cm(2.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
 s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '微软雅黑')
 s.font.size=Pt({'Title':26,'Subtitle':12,'Heading 1':18,'Heading 2':12,'Heading 3':11}.get(name,10.5))
 s.paragraph_format.space_after=Pt(7)
 s.paragraph_format.line_spacing=1.2
doc.styles['Normal'].paragraph_format.widow_control=True
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.styles['Subtitle'].font.italic=False
doc.core_properties.title='晶织 多服务器连接与计算控制智能体作品设计文档'
doc.core_properties.subject='一〇七杯算力与智能体开发大赛'
doc.core_properties.author=''
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=foot.add_run(); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); r._r.addnext(f)

def p(t): return doc.add_paragraph(t)
def h(t): return doc.add_heading(t,2)
def page(t): doc.add_page_break(); doc.add_heading(t,1)
def table(head,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(head)); t.style='Table Grid'; t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths): c.width=Cm(w)
 for c,v in zip(t.rows[0].cells,head): c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row): c.text=v
 for i,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   if widths: c.width=Cm(widths[j])
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(5); pp.paragraph_format.space_before=Pt(5)
    pp.paragraph_format.line_spacing=1.1
    for r in pp.runs: r.font.size=Pt(9.5); r.bold=(i==0)
   if i==0:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EEF1F4'); c._tc.get_or_add_tcPr().append(sh)
  tr=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(tr)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 p('')
 return t

# A code-drawn architecture diagram with Chinese labels.
im=Image.new('RGB',(1500,1050),'white'); d=ImageDraw.Draw(im)
font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',29)
small=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',25)
def box(x,y,w,hh,title,body):
 d.rounded_rectangle((x,y,x+w,y+hh),radius=6,fill='#F4F6F8',outline='#425466',width=3)
 d.text((x+22,y+16),title,font=font,fill='black')
 for n,line in enumerate(body): d.text((x+22,y+61+n*35),line,font=small,fill='#263747')
def arrow(x,y,xx,yy):
 d.line((x,y,xx,yy),fill='#425466',width=4)
 d.polygon([(xx,yy),(xx-9,yy-16),(xx+9,yy-16)],fill='#425466')
box(70,20,1360,120,'用户入口',['本地 Web 控制台    CLI    MCP    VS Code Bridge'])
arrow(750,140,750,180)
box(70,180,840,170,'本地智能体与应用服务',['Python 标准库 HTTP 服务  ·  对话与项目状态','Provider 适配 → 工具调用循环 → ToolRegistry'])
box(960,180,470,170,'模型服务  仅本机访问',['Chat 兼容 / Responses / Codex','经配置的 API 或本地桥接调用'])
d.line((910,260,960,260),fill='#425466',width=4)
arrow(490,350,490,390)
box(70,390,840,170,'标准单阶段执行路径',['输入校验 → 不可变计划 → 人工审批 → 执行引擎','本地记录：计划  ·  审批账本  ·  作业状态  ·  审计'])
box(960,390,470,170,'高级命令入口',['shell_run / remote_run','审计型工具，独立于标准计划门禁'])
arrow(490,560,490,605)
box(70,605,1360,145,'多服务器连接与路由  SSH 与 Gateway',['本机 → Vlab 中转 → 多台服务器，或本机分别 SSH 直连','命名操作执行参数与路径校验；高级命令沿独立路径执行'])
arrow(750,750,750,795)
box(70,795,430,145,'服务器 A  Slurm',['队列与资源查询','VASP 专用计算流程'])
box(535,795,430,145,'服务器 B  PBS',['文件与作业管理','按目标配置独立执行'])
box(1000,795,430,145,'服务器 C  SSH',['目录浏览与文件传输','返回实际状态和日志'])
d.text((90,955),'计算服务器无需访问模型 API；模型调用与结果解释在本机侧完成',font=small,fill='#263747')
d.text((70,982),'边界：大模型负责规划与解释；VASP 计算在集群资源上执行。',font=small,fill='#263747')
im.save(ROOT/'architecture.png')

doc.add_paragraph('晶织 多服务器连接与计算控制智能体',style='Title')
doc.add_paragraph('作品设计文档',style='Title')
doc.add_paragraph('一〇七杯算力与智能体开发大赛',style='Subtitle')
p('实现标识  rsci    文档版本  1.1    修订日期  2026年9月6日')
h('作品概述')
p('晶织面向科研服务器无法直接访问外部模型服务、且用户需要管理多台服务器的场景，提供本地智能体驱动的统一连接与计算控制工作台。模型调用在可联网的本机完成，服务器通过 SSH 直连或 Vlab 中转接收已授权操作，无需直接调用模型 API。用户可集中诊断连接、管理文件、查询作业并解释日志。')
p('设计主线是把本机的智能交互与多台远端服务器的真实执行环境连接起来，减少切换终端、搬运命令与日志以及混淆目标目录的操作。VASP 是当前专业适配场景，已有输入校验、计划审批、单阶段提交、监测与结果回收；通用连接层与科学软件专用逻辑分开组织。')
h('阅读导航')
table(['章节','内容'],[
 ('1  需求与设计思路','真实场景、目标、范围与设计取舍'),
 ('2  技术架构','总体架构图、部署边界与数据流'),
 ('3  功能模块','交互、文件、作业、监测和扩展接口'),
 ('4  智能体与模型调用','协议适配、工具循环、能力探测和异常处理'),
 ('5  工作流与数据设计','多服务器与输入门禁'),
 ('6  安全设计与技术难点','凭据、执行边界、并发与科学判读'),
 ('7  部署与验证设计','运行环境、验收矩阵与算力证据'),
 ('8  创新价值与演进','创新点、完成范围及后续计划'),
 ('附录  实现依据','代码索引与提交配置记录')],[4.7,11.7])

page('1  需求与设计思路')
h('1.1  受限网络与多服务器操作痛点')
p('科研服务器可能禁止访问公网模型 API，或因软件安装权限、系统版本等限制难以部署智能体客户端。用户因而在本机向模型提问，在远端手工执行，再将日志复制回来解释。这种割裂使建议难以对应真实路径、环境和执行状态，也增加了复制命令出错的机会。')
p('多台服务器还具有不同地址、认证方式、工作根目录和调度器。用户在多个终端之间核对文件与作业，容易选错目标、遗漏失败信息或混用结果。系统先解决统一连接、上下文关联与状态集中查看，再以 VASP 专用解析补充科学收敛判断。')
h('1.2  用户与需求')
table(['用户','主要需求','对应设计'],[
 ('科研初学者','理解文件、提交过程和错误原因','自然语言入口、输入检查、计划卡片'),
 ('研究人员','跨服务器管理任务，避免重复劳动','统一服务器目录、作业台账、结果回收'),
 ('指导教师与维护人员','了解执行内容和失败原因','文件哈希、执行尝试、审计及诊断')],[3,5.6,7.8])
h('1.3  设计目标与边界')
p('首要目标是在同一本地入口登记、连接和操作多台授权服务器，让每次文件或作业操作绑定明确目标。其次在既有连接上形成可检查的专业计算流程。标准流程对输入变更重新审批，对提交结果不确定的任务停止自动重提；界面展示真实工具结果和失败位置。')
p('系统复用现有 SSH、Slurm/PBS 与服务器软件环境，不另建调度器。当前没有全局最优调度或自动跨服务器迁移能力；多服务器管理指统一接入和分别操作。VASP 专用流程已经实现，其他科学软件仍需补充参数校验、输出解析和工作流适配。')
h('1.4  主要设计取舍')
p('模型接入放在能够访问所配置服务的本机，远端只运行传输、网关与计算所需组件，降低对计算服务器外网和软件环境的依赖。本机仍须能访问模型服务并通过 SSH 到达目标；该设计不绕过网络策略，也不意味着完全离线运行。发给远程模型的必要文本仍会离开本机。')

page('2  技术架构')
p('系统分为本地交互与智能控制、连接与路由、多服务器执行三层。只有本机 Provider 调用模型服务，SSH 层按服务器配置直连或经 Vlab 中转。不同服务器分别承担文件操作和计算任务，计算端不承担模型 API 访问。')
doc.add_picture(str(ROOT/'architecture.png'),width=Cm(14.5))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
p('图 1  本机模型接入与多服务器执行架构')
h('2.1  实现技术与部署位置')
table(['层级','技术实现','职责'],[
 ('本地交互','原生 HTML / CSS / JavaScript','浏览器控制台与流式工具反馈'),
 ('应用与控制','Python 3.11+，ThreadingHTTPServer','会话、模型调用、项目、计划和审批'),
 ('远程连接','OpenSSH、Python Gateway','Vlab 中转或目标 HPC 直连'),
 ('计算执行','Slurm / PBS、MPI、VASP','资源排队、作业执行、结果输出'),
 ('辅助入口','CLI、MCP、TypeScript 扩展','复用能力并提供 VS Code 文件访问')],[3.1,6.1,7.2])
h('2.2  数据流与控制流')
p('用户选择服务器和目录后，本机将必要上下文发送给模型；模型返回工具名称与参数，由注册表调用连接层。每台服务器按自己的路由、认证和根路径执行，作业状态与日志返回本机供界面和模型解释。VASP 标准计划在本地审批后执行；高级 Shell 的独立边界见第 6 章。')

page('3  功能模块')
table(['模块','主要功能','输入与输出'],[
 ('智能对话','流式回答、工具调用卡片、会话持久化、分支与压缩','用户目标与上下文 → 建议、工具结果或待审批计划'),
 ('项目与输入','项目目录管理、计算参数编辑、文件读写和校验','真实结构与参数 → 本地 VASP 输入文件'),
 ('服务器与连接','服务器登记、Vlab/直连路由、诊断和凭据管理','连接配置 → 可用状态与错误信息'),
 ('远端文件','目录浏览、小文本读取、上传下载、回收站','明确路径 → 文件元数据、内容或传输结果'),
 ('计划中心','预览步骤、文件哈希、审批、执行状态和恢复','不可变计划与人工审批 → 持久化运行记录'),
 ('作业与科学进度','活动及历史作业、目录关联、进度与规则诊断','调度器及 VASP 输出 → 分层状态与诊断依据'),
 ('资源监测','节点负载、内存和可用 GPU 指标、队列概况','带采集时间的快照 → 资源视图'),
 ('编辑与扩展','VS Code 虚拟目录、安全编辑及可选 Vlab 工作区','明确选中的计算目录 → 编辑访问与同步结果')],[3,6.7,6.7])
h('3.1  从多服务器连接进入实际操作')
p('用户先登记多台服务器的地址、认证信息、根目录和路由，在统一入口诊断连接、浏览目录、查看队列和资源快照，再指定目标开展操作。VASP 任务可进一步生成计划卡片，固定服务器、输入、脚本和资源参数，审批后提交；缺少真实结构或关键参数时先补齐资料。')
h('3.2  从监测返回可解释结果')
p('作业页面将作业编号和计算目录关联，避免只显示编号而无法找到结果。进度读取使用受限日志和专用解析器；界面将调度器状态与科学状态分别展示。缺失日志或采集失败应体现为未知或失败信息，不能用空值代替“正常”。资源指标用于判断平台负载，不等同于调度器的资源预留承诺。')
h('3.3  编辑与检索的定位')
p('VS Code Bridge 经本地服务访问计算文件，适合不能安装现代 VS Code Server 的老集群。完整工作区模式依赖 Vlab 侧额外部署。网页检索属于可选辅助信息来源，不能直接替代运行日志，也不意味着系统已经建设了向量检索知识库或 RAG 平台。')

page('4  智能体与模型调用')
h('4.1  模型接入方式')
p('模型服务由 Provider 配置选择，包含 id、protocol、base_url、model 和 api_key_env。系统不把单一模型型号固化为执行依赖；更换模型时，工具接口和工作流引擎保持一致。参赛演示所用的具体模型应与实际部署配置一致，记录项见附录。')
table(['协议标识','调用方式','适配职责'],[
 ('openai-chat-compatible','POST 到配置地址的 /chat/completions','messages、tools、SSE 流式文本与函数参数'),
 ('openai-responses','调用配置的 Responses 接口','将响应事件和函数调用归一化'),
 ('codex-sdk','经本地 Node 桥接 SDK 或 CLI','归一化模型输出，工具分派留在 Python 层')],[5,5.4,6])
h('4.2  工具调用循环')
p('① 组装系统约束、用户目标和历史上下文。② Provider 发送请求并接收文本或工具调用。③ 注册表按名称选择工具，检查 Provider 模式，由具体处理器校验参数并执行。④ 工具结果经大小限制和审计处理后回传模型及界面。⑤ 模型继续解释或调用下一工具，达到轮次上限或收到停止信号则终止；产生标准计划审批卡时暂停。')
p('系统提示明确要求优先使用命名工具、读取现有文件、避免编造结构、区分科学收敛与调度状态。提示约束用于指导模型行为；涉及可靠执行的控制仍由计划哈希、审批校验和执行状态落实。提示文本本身不构成对高级 Shell 命令的强制隔离。')
h('4.3  能力探测与故障处理')
p('Provider 检测可达性、流式输出和工具调用能力，三项满足时进入 full 模式；否则降为 analysis_only，注册表拒绝 write 工具。structured_json 作为探测信息保留，当前代码不把它列为 full 模式的必要条件。')
p('Chat 兼容协议使用 temperature=0 降低随机波动，但不能保证回答确定或正确。HTTP 传输对配置范围内的临时错误、限流及服务故障进行有限重试，默认最多额外重试 3 次，并采用有上限的退避。网络重试与集群提交恢复分离，不能因模型连接失败而直接重新提交作业。')
h('4.4  上下文与凭据')
p('会话支持持久化、分支和自动压缩，便于持续处理科研项目；压缩摘要不能替代执行计划与审计记录。API Key 优先从环境变量读取，也支持从本机 DPAPI 加密存储解析。模型仅应获得任务必要的科研文本，POTCAR 在专用文件接口中只处理元数据。')

page('5  工作流与数据设计')
h('5.1  单阶段执行过程')
p('准备输入 → 本地校验 → 生成计划 → 人工审批 → mkdir → upload → validate → submit → monitor → progress → download → parse。每次计划对应一个阶段、一个目标服务器和一套确定的输入；relax、static、DOS 等阶段需要分别准备与审批。')
p('本地和远端门禁检查 INCAR 的有效参数、KPOINTS 格式、POSCAR 晶格与坐标完整性，以及 POTCAR 是否存在。skip-potcar 仅表示不上传本地 POTCAR，远端仍必须预先具备正确的已组装 POTCAR；当前单阶段引擎不会自动从库中拼接赝势。')
h('5.2  关键数据对象')
table(['对象','主要内容','设计作用'],[
 ('服务器配置','标识、地址、根路径、路由、调度器','固定命名操作的目标与目录范围'),
 ('执行计划','plan_id、plan_hash、文件 SHA-256、脚本、步骤与资源参数','把审批对象绑定到实际执行内容'),
 ('审批令牌','server、plan_hash、files_hash、action、parameter_hash、有效期、签名','拒绝过期、变更及跨用途重放'),
 ('运行与作业记录','尝试、步骤状态、job_id、目录、调度及科学状态','支持查询、失败定位和恢复判断'),
 ('会话与审计','对话上下文、工具事件、结果摘要','支持解释、排错与过程追踪')],[3.1,7.3,6])
h('5.3  审批一致性与持久化')
p('计划根据服务器、输入文件、脚本和参数生成确定性哈希，审批令牌以本地 HMAC 签名绑定这些内容。标准路径中，输入或脚本改变后需要重建计划并重新审批；上传前再次校验本地 SHA-256，防止审批后文件被悄然替换。')
p('审批消费和运行状态写入使用本地锁；关键状态采用临时文件、fsync 与原子替换保存。每个外部步骤执行前记录 started。若提交阶段发生进程中断且不能确认是否已获得作业编号，恢复进入 needs_review，先人工核对调度器，避免盲目再次 sbatch 或 qsub。')
h('5.4  结果与状态语义')
p('调度器 COMPLETED 仅说明调度生命周期结束。工作流 completed 还需要科学解析确认相应判据满足；结束但未收敛进入 needs_review，明确失败进入 failed。结果解析连同 INCAR 读取，避免把缺失 NSW 的任务误判为静态计算；静态任务缺少可选 CONTCAR 不应导致整个回收失败。已有本地结果与新下载副本哈希不一致时拒绝混用。')

page('6  安全设计与技术难点')
h('6.1  安全控制的适用范围')
p('本地控制台默认绑定 127.0.0.1，API 请求使用启动时生成的会话令牌。SSH 连接校验主机身份；服务器密码和 TOTP 种子可通过 Windows DPAPI CurrentUser 加密保存，按确认周期控制自动使用。Vlab 中转时发送当前认证所需信息，TOTP 种子保留在本机。')
p('命名文件工具实施路径和参数校验，标准工作流使用不可变计划及审批门禁。但当前 full 模式还提供 shell_run 与 remote_run，两者为审计型通用命令入口，未逐条拦截命令或要求标准计划审批。因此，不能把标准路径的目录限制、POTCAR 限制和审批保护描述成覆盖全部入口的强制安全沙箱。')
p('当前部署适用于受信任操作者管理的个人科研环境。若面向不受信任用户或开放多租户服务，应先收紧通用命令入口，建立统一强制授权，并验证隔离边界；这属于后续加固工作。审计能辅助追踪，但不等于命令执行前阻断，也不提供独立防篡改证明。')
h('6.2  技术难点与处理方式')
table(['难点','设计处理','仍需关注的边界'],[
 ('跨跳认证与连接差异','Vlab/直连共用客户端接口，本机加密凭据，连接诊断','依赖网络、主机指纹与集群认证策略'),
 ('自然语言转实际操作','统一工具注册表、能力探测、标准计划交接','通用 Shell 仍具有较宽执行能力'),
 ('审批后内容变化','计划及文件哈希绑定，上传前复核','标准路径外操作需单独治理'),
 ('重复提交与异常中断','外部步骤前落盘，提交不确定时转人工核对','不能对所有远端故障宣称严格恰好一次'),
 ('调度状态与科学判据错位','专用 VASP 解析，分开展示生命周期与收敛','有限日志和复杂计算仍需科研人员复核'),
 ('轻量部署与生态扩展','默认标准库执行，科学包按需安装','可选依赖安装不代表自动化链已接通')],[3.5,6.7,6.2])
h('6.3  性能与资源使用原则')
p('文件查看采用大小与类型限制，模型上下文采用摘要与有界工具结果，避免反复读取大型科学二进制文件。资源监测与作业查询使用快照，展示采集时间以识别陈旧信息。效率评估关注人工操作步骤、提交成功率、错误定位耗时和模型响应延迟；VASP 运行时间主要取决于计算参数、硬件与集群配置，不能把界面自动化等同于求解加速。')

page('7  部署与验证设计')
h('7.1  部署条件与启动流程')
p('本机需要 Python 3.11 或更高版本、OpenSSH 和浏览器。集群需具备 Python 3、Slurm 或 PBS，以及有权使用的 VASP、MPI 和赝势环境。Vlab 路由部署 Gateway；直连路由使用目标端辅助程序。VS Code Bridge 和完整工作区按需安装，不是核心提交流程的前置条件。')
p('建议依次完成：安装项目 → 启动 rsci ui → 注册服务器和工作根目录 → 完成认证 → 只读诊断连接与队列 → 配置模型并探测 → 校验小型算例 → 审阅并执行标准计划 → 回收结果。首次验证使用小规模任务，明确分区、核数、时限和 VASP 模块。')
h('7.2  验收矩阵')
p('以下为验证方案与通过判据，不作为本次已取得的实测结果。仓库已有相关自动化测试和 Slurm 算例；正式交付时应在确定的提交版本上保存测试记录。')
table(['验证项','方法与预期结果'],[
 ('多服务器与输入门禁','登记两台服务器并分别查询文件与队列，核对目标对应；VASP 非法输入应阻止标准流程继续'),
 ('审批与重放','改变输入使旧审批失效；过期或跨运行重放被拒绝'),
 ('中断恢复','模拟提交时断开，转 needs_review，不能自动重复提交'),
 ('科学状态','分别用静态、弛豫、未收敛及缺失日志验证解析语义'),
 ('Provider 降级','探测失败后 write 工具被拒绝，仍可执行只读任务'),
 ('真实集群闭环','一次受审阅的小型任务完成提交、监测、回收及目录对应'),
 ('受限网络与部署','记录计算服务器无法访问模型 API 的环境，验证本机可调用模型且 SSH 操作往返成功')],[3.5,12.9])
h('7.3  算力运行证据采集')
p('若参加算力平台赛道，需另附本科生算力平台 Slurm 集群的真实运行说明。仓库 examples/slurm-evidence 提供 Si 静态、Al 弛豫和 Fe 自旋弛豫小型算例，以及 collect_evidence.sh。算例参数用于功能验证，不能直接作为已完成收敛测试的科研参数。')
p('每次运行记录作业编号、提交与启动时间、排队时长、运行时长、分区、AllocCPUS、MaxRSS、ExitCode、软件环境、输入哈希和输出摘要。内存数据取自平台实际记账字段，缺失应注明不可用；不同任务步的数据需标明口径。采集脚本只对已有任务收集证据，不替代实际运行。')
p('对照实验可在相同输入和资源条件下比较手工流程与智能体流程，统计人工交互次数、准备耗时与差错情况；另行报告模型响应耗时，避免与队列等待、VASP 计算时长混为一项。')

page('8  创新价值与演进')
h('8.1  创新点与实际价值')
p('第一，将智能体与计算服务器分开部署。本机负责访问模型，连接层把工具操作送达受限网络中的服务器，再回传真实执行结果，使智能交互能够参与现有科研环境，减少手工搬运命令和日志。')
p('第二，以明确的服务器和目录关联文件、队列、资源与任务记录。多服务器使用者能够从同一入口查看状态、选择目标并追踪执行。标准 VASP 流程进一步通过不可变计划绑定文件、资源参数和审批内容，降低操作目标与执行内容不一致的风险。')
p('第三，将通用连接控制与专业计算适配分层。以 VASP 为落地场景，实现从输入检查到结果判读的闭环，并区分调度完成与科学收敛。其他软件可复用连接基础，但需单独验证专业逻辑；上述方案特点不宣称同类产品中首次实现。')
h('8.2  当前实现范围')
table(['能力','当前状态'],[
 ('本地控制台及多入口','已有 Web、CLI、MCP 与 VS Code Bridge 实现'),
 ('模型接入与工具循环','已有三种协议适配、流式输出与能力模式控制'),
 ('单阶段工作流','已有准备、审批、执行、状态持久化与恢复逻辑'),
 ('远端与科学计算','已有 Vlab/直连、Slurm/PBS 适配、VASP 检查与解析'),
 ('材料库增强','pymatgen 可选接入局部结构与输入校验'),
 ('多阶段与流程生态','custodian、atomate2、DPDispatcher、AiiDA 尚非默认执行链'),
 ('统一强制隔离','标准路径已有门禁，通用 Shell 入口仍需进一步加固')],[5,11.4])
h('8.3  后续演进顺序')
p('优先验证多服务器连接稳定性、路由诊断、断连反馈和操作目标一致性，补齐真实运行证据并加固权限边界。随后完善统一作业与文件视图，并按实际需求增加科学软件适配；分别验证其输入和结果语义。')
p('有明确需求后再引入多阶段编排：用 pymatgen 加强材料数据处理，以 custodian 表达受审批的修复策略，以 atomate2 描述阶段依赖，以 DPDispatcher 适配执行后端；只有跨服务器持久流程和数据溯源需求达到相应规模时，再评估 AiiDA。每一步都应提供与现有 Gateway、审批和状态模型之间清晰的职责边界。')

page('附录  实现依据与配置记录')
h('A  设计依据')
p('本设计文档于2026年9月6日依据当前工作区修订，重点说明受限网络下的多服务器连接与控制。赛事提交要求沿用原稿所列预通知；实际赛道、模型配置和真实运行记录应在提交前核实。')
table(['实现主题','仓库内依据'],[
 ('入口与界面','src/rsci/ui/server.py；ui/static/index.html'),
 ('模型与智能体','src/rsci/providers/；agents/runtime.py；agents/memory.py'),
 ('工具与权限模式','src/rsci/tools/registry.py'),
 ('计划及审批','src/rsci/workflow/plan.py；approval.py；engine.py'),
 ('连接及凭据','src/rsci/gateway/；core/config.py；core/secretbox.py'),
 ('调度与科学解析','src/rsci/hpc/scheduler.py；jobscript.py；vasp.py'),
 ('可选材料集成','src/rsci/integrations/materials.py；spec/automation-stack.md'),
 ('测试及证据准备','tests/；examples/slurm-evidence/'),
 ('工作流范围','spec/workflows.md；recipes/README.md')],[4.4,12])
h('B  参赛演示配置记录')
p('此表用于固定演示与部署程序的一致性，不填写密钥、口令、TOTP 种子或私钥正文。')
table(['项目','记录内容'],[
 ('队伍名称与赛道','待填写'),
 ('实际模型及服务提供方','待填写具体型号、提供方及版本信息'),
 ('调用协议及服务地址','待填写 protocol、base_url；不包含 API Key'),
 ('部署版本','待填写提交包版本及 Git 提交标识'),
 ('真实运行环境','待填写集群、分区、VASP/MPI 版本及作业编号')],[5,11.4])
p('提交前需核对程序版本元数据、演示配置与测试记录的一致性。本设计文档修订版本为 1.1；文中验收矩阵是验证方案，不代表所有项目均已完成实测。')

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
