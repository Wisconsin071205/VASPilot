from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent/'competition_submission'/'修订版'/'晶织_作品简介.docx'
doc=Document()
s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(2.1);s.bottom_margin=Cm(2.1)
s.left_margin=Cm(2.3);s.right_margin=Cm(2.3)
for name,size in [('Normal',11),('Title',24),('Subtitle',11),('Heading 1',15),('Heading 2',12)]:
 st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size)
 st.font.color.rgb=RGBColor(0,0,0);st.font.italic=False
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '微软雅黑')
 st.paragraph_format.space_after=Pt(8);st.paragraph_format.line_spacing=1.2
for st in doc.styles:
 for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
doc.core_properties.title='晶织 多服务器连接与计算控制智能体作品简介'
doc.core_properties.subject='一〇七杯算力与智能体开发大赛'
doc.core_properties.author=''
ft=s.footer.paragraphs[0];ft.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');ft._p.append(fld)
def p(t):doc.add_paragraph(t)
def h(t):doc.add_heading(t,1)
def feature(title,body):
 pp=doc.add_paragraph();pp.add_run(title+'  ').bold=True;pp.add_run(body)

doc.add_paragraph('晶织 多服务器连接与计算控制智能体',style='Title')
doc.add_paragraph('作品简介',style='Subtitle')
p('一〇七杯算力与智能体开发大赛  |  技术名称 rsci')
p('晶织面向无法直接访问外部模型服务、且需要同时管理多台科研服务器的使用场景。智能体运行在可访问模型服务的本机，通过 SSH 直连或 Vlab 中转连接服务器，在一个工作台中完成连接诊断、文件管理、作业查询与结果解释。计算服务器无需直接访问模型 API；VASP 是当前已实现专用工作流的应用场景。')
h('1  作品背景与解决的问题')
p('科研服务器常受外网访问、软件安装和运行环境限制，无法像个人电脑一样直接接入在线智能体。用户只能在本机询问模型，再把命令和文件搬到服务器执行，反复复制日志反馈，智能建议与真实环境之间容易脱节。')
p('当任务分布在多台服务器时，地址、认证方式、工作目录和调度器各不相同，用户还需反复切换终端，核对文件版本和作业位置。晶织把模型调用留在本机，将已授权的远端操作交给连接层执行，让用户围绕明确的服务器、目录和任务持续工作。')
h('2  核心功能')
feature('本机智能体与远端协作','模型在本机侧接收任务和必要上下文，工具读取实际文件与状态，再把结果反馈给用户；支持流式对话、会话分支与上下文压缩。')
feature('多服务器统一连接','登记多台服务器，分别配置认证、根目录和连接路由；支持 Vlab 中转或直接 SSH，集中查看连接状态、远端文件、Slurm/PBS 作业及资源快照。')
feature('文件与任务集中管理','浏览、编辑和传输远端文件，关联作业编号与计算目录。当前 VASP 标准流程还提供输入校验、计划审批、提交、监测和结果回收。')
feature('过程可核对','标准计划展示服务器、资源、脚本与文件哈希；提交状态不确定时转人工核对。VASP 专用解析区分调度完成与科学收敛，并保留日志依据。')
h('3  典型使用场景')
p('研究人员在本机打开工作台，查看两台已登记服务器的连接、队列和资源快照，选定目标服务器后浏览计算目录、检查文件并查询作业。若任务使用 VASP，可继续生成计划、人工审批并提交，随后在同一入口观察进度和回收结果。目标选择由用户或明确任务参数确定，不宣称自动跨服务器迁移作业。')

h('4  模型及 API 调用方式')
doc.paragraphs[-1].paragraph_format.page_break_before = True
p('模型请求由本机发出，远端服务器负责执行授权操作和数值计算。本机需要能访问所配置的模型服务，并能直接或经中转到达服务器；“服务器无需联网调用模型”不等于整套系统完全离线，也不绕过集群网络和权限策略。')
t=doc.add_table(rows=1,cols=2);t.style='Table Grid';t.autofit=False
rows=[('接入协议','调用方式'),('OpenAI Chat 兼容协议','向配置的 /chat/completions 接口发送 messages 和 tools，使用 SSE 接收流式文本及工具调用。'),('Responses 协议','调用配置的 Responses 接口，将响应事件与函数调用转换为统一内部格式。'),('Codex SDK 或 CLI','通过本地 Node 桥接接收模型输出，由 Python 注册表分派工具。')]
for i,row in enumerate(rows):
 cells=t.rows[0].cells if i==0 else t.add_row().cells
 for j,txt in enumerate(row):
  cells[j].text=txt;cells[j].width=Cm([5,11.4][j])
  for pp in cells[j].paragraphs:
   pp.paragraph_format.space_after=Pt(5);pp.paragraph_format.space_before=Pt(5)
   for r in pp.runs:r.font.size=Pt(10);r.bold=i==0
 for cell in cells:
  if i==0:
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEF1F4');cell._tc.get_or_add_tcPr().append(sh)
p('可通过 OpenAI Chat 兼容协议配置 DeepSeek 等模型服务；具体提供方、型号和地址以实际演示配置为准，不在材料中预设未核实的模型版本。模型返回工具调用后，由本地服务执行并回传结果。')
p('API Key 从环境变量或本机 DPAPI 加密存储读取，不写入参赛材料。模型服务经过可达性、流式输出和工具调用能力探测；未满足执行条件时降级为只读分析。临时网络错误采用有限重试，标准工作流状态独立保存，避免把模型请求重试变成重复提交。')
h('5  创新点与实用价值')
feature('把智能体接到受限网络服务器','将模型访问与服务器执行分开部署，减少“本机问模型、远端抄命令、再复制日志”的往返操作。')
feature('面向多服务器组织操作','以服务器和目录标识关联连接、文件、作业与状态，减少多终端切换中的目标混淆；标准计算计划进一步绑定审批和执行内容。')
feature('以 VASP 验证专业任务闭环','在通用连接与管理能力上实现 VASP 输入检查和结果解释；其他软件的专用校验、解析和工作流需要独立适配。')
h('6  实现范围与推广方向')
p('当前已实现本地 Web 控制台、模型工具调用、多服务器登记、Vlab/直连、文件操作、Slurm/PBS 查询及单阶段 VASP 工作流。标准流程有审批门禁，高级 Shell 入口仅提供审计，不构成全面隔离的安全沙箱。')
p('作品适用于受信任用户管理的科研服务器与教学集群。可推广的是连接与控制架构；自动跨集群调度、完全离线模型推理和任意科学软件即插即用均不属于当前承诺。效率收益应通过真实任务测量。')
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
